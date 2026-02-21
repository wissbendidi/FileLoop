"""
Tests for the content extraction and chunking module.
"""

import os
import tempfile
import shutil
from pathlib import Path

import pytest

from src.content import (
    ContentExtractor,
    ContentChunk,
    TEXT_EXTENSIONS,
    PDF_EXTENSIONS,
    DOCX_EXTENSIONS,
)


class TestContentChunk:
    """Test ContentChunk dataclass."""

    def test_to_dict(self):
        chunk = ContentChunk(
            path="/foo/bar.txt",
            chunk_index=0,
            text="hello world",
            start_char=0,
            end_char=11,
            metadata={"page": 1},
        )
        d = chunk.to_dict()
        assert d["path"] == "/foo/bar.txt"
        assert d["chunk_index"] == 0
        assert d["text"] == "hello world"
        assert d["start_char"] == 0
        assert d["end_char"] == 11
        assert d["metadata"] == {"page": 1}


class TestContentExtractor:
    """Test ContentExtractor extraction and chunking."""

    def setup_method(self):
        self.temp_dir = tempfile.mkdtemp()
        self.extractor = ContentExtractor(
            chunk_size=200,
            chunk_overlap=20,
        )

    def teardown_method(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_extract_text_simple(self):
        path = os.path.join(self.temp_dir, "readme.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("Hello world.\nThis is a test file.")
        chunks = self.extractor.extract_file(path)
        assert len(chunks) == 1
        assert chunks[0].path == str(Path(path).resolve())
        assert chunks[0].chunk_index == 0
        assert "Hello world" in chunks[0].text
        assert "test file" in chunks[0].text

    def test_extract_text_chunking(self):
        path = os.path.join(self.temp_dir, "long.txt")
        # Create text longer than chunk_size (200) so we get multiple chunks
        lines = ["Paragraph " + str(i) + " content here.\n" for i in range(30)]
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n\n".join(lines))
        chunks = self.extractor.extract_file(path)
        assert len(chunks) >= 2
        all_text = " ".join(c.text for c in chunks)
        assert "Paragraph 0" in all_text
        assert "Paragraph 29" in all_text

    def test_extract_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            self.extractor.extract_file(os.path.join(self.temp_dir, "nonexistent.txt"))

    def test_extract_file_unsupported_extension(self):
        path = os.path.join(self.temp_dir, "image.png")
        with open(path, "wb") as f:
            f.write(b"fake png")
        with pytest.raises(ValueError, match="Unsupported extension"):
            self.extractor.extract_file(path)

    def test_extract_directory_empty(self):
        chunks = self.extractor.extract_directory(self.temp_dir)
        assert chunks == []

    def test_extract_directory_with_text_files(self):
        for name in ["a.txt", "b.md"]:
            with open(os.path.join(self.temp_dir, name), "w", encoding="utf-8") as f:
                f.write("Content of " + name)
        chunks = self.extractor.extract_directory(self.temp_dir)
        assert len(chunks) >= 2
        paths = {c.path for c in chunks}
        assert any("a.txt" in p for p in paths)
        assert any("b.md" in p for p in paths)

    def test_extract_directory_excludes_dirs(self):
        git_dir = os.path.join(self.temp_dir, ".git")
        os.makedirs(git_dir)
        with open(os.path.join(git_dir, "config.txt"), "w", encoding="utf-8") as f:
            f.write("git config")
        with open(os.path.join(self.temp_dir, "root.txt"), "w", encoding="utf-8") as f:
            f.write("root file")
        chunks = self.extractor.extract_directory(
            self.temp_dir, excluded_dirs=[".git"]
        )
        path_strs = [c.path for c in chunks]
        assert any("root.txt" in p for p in path_strs)
        assert not any(".git" in p for p in path_strs)

    def test_extract_pdf(self):
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            pytest.skip("PyPDF2 not installed")
        path = os.path.join(self.temp_dir, "doc.pdf")
        # Create a minimal PDF with PyPDF2
        from PyPDF2 import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with open(path, "wb") as f:
            writer.write(f)
        # Re-open and add minimal text (PyPDF2 doesn't make it trivial to add text)
        reader = PdfReader(path)
        page = reader.pages[0]
        text = page.extract_text() or ""
        # Blank page may have no text; at least extraction should not crash
        chunks = self.extractor.extract_file(path)
        # Either we get no chunks (blank) or some text
        assert isinstance(chunks, list)

    def test_extract_docx(self):
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")
        path = os.path.join(self.temp_dir, "doc.docx")
        doc = Document()
        doc.add_paragraph("First paragraph in DOCX.")
        doc.add_paragraph("Second paragraph.")
        doc.save(path)
        chunks = self.extractor.extract_file(path)
        assert len(chunks) >= 1
        full = " ".join(c.text for c in chunks)
        assert "First paragraph" in full
        assert "Second paragraph" in full

    def test_code_chunking_prefers_units(self):
        path = os.path.join(self.temp_dir, "code.py")
        content = '''
def foo():
    return 1

class Bar:
    def baz(self):
        return 2
'''
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        chunks = self.extractor.extract_file(path)
        assert len(chunks) >= 1
        all_text = " ".join(c.text for c in chunks)
        assert "def foo" in all_text or "foo" in all_text
        assert "class Bar" in all_text or "Bar" in all_text
